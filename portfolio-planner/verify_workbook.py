#!/usr/bin/env python3
"""
Acceptance tests for Portfolio_Planner_CA.xlsx (spec §3 + §7).

1. Recalculates the workbook headlessly with LibreOffice (openpyxl writes no
   cached values, so every number read back was computed from the formulas).
2. Compares engine, tax, projection and dashboard values against the
   independent Python reference model (reference_model.py).
3. Checks the v1 bug list: withdrawals flow through from month 61, realized
   gains appear when dividends fall short, Withdrawals Summary populates,
   Dashboard KPIs tie to the Engine, and no error cells outside ChartData.
"""

import subprocess, sys, os, shutil, tempfile
import openpyxl
from reference_model import Inputs, simulate, annual_tax

SRC = os.path.join(os.path.dirname(__file__), "Portfolio_Planner_CA.xlsx")
ROW0, FIRST = 3, 4

fails = []


def check(name, ok, detail=""):
    print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"  [{detail}]" if detail and not ok else ""))
    if not ok:
        fails.append(name)


def close(a, b, tol=0.01, rel=1e-6):
    if a is None or b is None:
        return False
    return abs(a - b) <= max(tol, rel * max(abs(a), abs(b)))


RECALC_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema">
 <item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
 <item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
 <item oor:path="/org.openoffice.Office.Calc/Formula/Calculation"><prop oor:name="UseThreadedCalculationForFormulaGroups" oor:op="fuse"><value>false</value></prop></item>
</oor:items>
"""
# LibreOffice's threaded formula-group calculation can mis-order same-row
# cross-column dependencies (harmless in Excel/Google Sheets, which walk the
# full dependency graph); disable it for deterministic verification.


def recalc(src):
    tmp = tempfile.mkdtemp()
    prof = os.path.join(tmp, ".config", "libreoffice", "4", "user")
    os.makedirs(prof)
    with open(os.path.join(prof, "registrymodifications.xcu"), "w") as f:
        f.write(RECALC_XCU)
    shutil.copy(src, os.path.join(tmp, "src.xlsx"))
    outdir = os.path.join(tmp, "out")
    os.makedirs(outdir)
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "xlsx",
         "--outdir", outdir, os.path.join(tmp, "src.xlsx")],
        check=True, capture_output=True, timeout=900,
        env={**os.environ, "HOME": tmp, "SC_NO_THREADED_CALCULATION": "1"},
    )
    return openpyxl.load_workbook(os.path.join(outdir, "src.xlsx"), data_only=True)


def col_letter_map():
    sys.path.insert(0, os.path.dirname(__file__))
    from build_workbook import EL
    return EL


def main():
    print("recalculating with LibreOffice…")
    wb = recalc(SRC)
    EL = col_letter_map()
    eng = wb["Engine (Monthly)"]
    nowd = wb["Engine (No Withdrawals)"]
    em1, ep1 = wb["Engine (-1%)"], wb["Engine (+1%)"]
    proj, tax, dash = wb["Projection"], wb["Tax Estimate"], wb["Dashboard"]

    def E(colname, month, sheet=eng):
        return sheet[f"{EL[colname]}{month + ROW0}"].value

    print("running reference model…")
    inp = Inputs()
    months = simulate(inp)
    taxes = annual_tax(months, inp)

    m = lambda t: months[t - 1]

    # --- engine vs reference model
    for t in [1, 12, 60, 61, 62, 96, 180, 240]:
        check(f"engine balance month {t} matches reference",
              close(E("BALTOT", t), m(t)["baltot"]),
              f"xlsx={E('BALTOT', t)} ref={m(t)['baltot']}")
    check("no withdrawal before start month (month 60)", (E("DRAWTOT", 60) or 0) < 1e-9)
    check("withdrawal target month 61 = $2,000", close(E("WT", 61), 2000))
    for t in [61, 100, 150, 240]:
        check(f"withdrawal drawn month {t} matches reference",
              close(E("DRAWTOT", t), m(t)["drawtot"]),
              f"xlsx={E('DRAWTOT', t)} ref={m(t)['drawtot']}")
    check("withdrawals indexed (month 73 target = 2000×1.02)", close(E("WT", 73), 2000 * 1.02))
    # forced sales & realized gains once dividends fall short
    ref_sale_y6 = sum(mm["saletot"] for mm in months if mm["year"] == 6)
    xl_sale_y6 = sum(E("SALETOT", t) or 0 for t in range(61, 73))
    check("forced sales occur in year 6 (dividends insufficient)", ref_sale_y6 > 0 and xl_sale_y6 > 0)
    check("year-6 sales match reference", close(xl_sale_y6, ref_sale_y6))
    ref_gain_y6 = sum(mm["gain"] for mm in months if mm["year"] == 6)
    xl_gain_y6 = sum(E("GAIN", t) or 0 for t in range(61, 73))
    check("realized capital gains in year 6 > 0", ref_gain_y6 > 0 and xl_gain_y6 > 0)
    check("year-6 realized gains match reference", close(xl_gain_y6, ref_gain_y6))
    check("balances reduced vs no-withdrawal engine",
          E("BALTOT", 240, nowd) - E("BALTOT", 240) > 100000)
    # frozen beyond horizon
    check("engine frozen after horizon", close(E("BALTOT", 300), E("BALTOT", 240)))

    # --- scenario engines
    m1 = simulate(Inputs(delta=-0.01))
    p1 = simulate(Inputs(delta=+0.01))
    check("scenario -1% ending balance matches reference", close(E("BALTOT", 240, em1), m1[239]["baltot"]))
    check("scenario +1% ending balance matches reference", close(E("BALTOT", 240, ep1), p1[239]["baltot"]))
    check("scenario ordering -1% < base < +1%",
          E("BALTOT", 240, em1) < E("BALTOT", 240) < E("BALTOT", 240, ep1))

    # --- tax tab vs reference
    for y in [1, 3, 6, 8, 15, 20]:
        r = y + 3
        check(f"taxable income year {y} matches reference",
              close(tax[f"K{r}"].value, taxes[y - 1]["taxable"]),
              f"xlsx={tax[f'K{r}'].value} ref={taxes[y - 1]['taxable']}")
        check(f"standalone tax year {y} matches reference",
              close(tax[f"S{r}"].value, taxes[y - 1]["standalone"]),
              f"xlsx={tax[f'S{r}'].value} ref={taxes[y - 1]['standalone']}")
        check(f"incremental (stacked) tax year {y} matches reference",
              close(tax[f"AD{r}"].value, taxes[y - 1]["incremental"]),
              f"xlsx={tax[f'AD{r}'].value} ref={taxes[y - 1]['incremental']}")
        # attribution layers must sum to the incremental (stacked) total
        layers = sum(tax[f"{L}{r}"].value or 0 for L in ("AN", "AO", "AP", "AQ"))
        check(f"tax-by-source layers sum to incremental total, year {y}", close(layers, tax[f"AD{r}"].value))

    # --- projection ties to engine (years 1, 8, final) + withdrawals summary
    for y in [1, 8, 20]:
        r = y + 3
        check(f"projection closing year {y} = engine month {y*12}",
              close(proj[f"AA{r}"].value, E("BALTOT", y * 12)))
    check("withdrawals summary populated (cum withdrawn year 8 > 0)",
          (proj["AB11"].value or 0) > 0)
    ref_cum_wd_y8 = sum(mm["drawtot"] for mm in months if mm["year"] <= 8)
    check("cum withdrawn year 8 matches reference", close(proj["AB11"].value, ref_cum_wd_y8))
    ref_cum_sale_y8 = sum(mm["saletot"] for mm in months if mm["year"] <= 8)
    check("cum funded-by-sales year 8 matches reference", close(proj["AD11"].value, ref_cum_sale_y8))

    # --- dashboard KPIs tie to engine (selected year = 8 by default)
    check("dashboard selected-year balance = engine month 96",
          close(dash["B9"].value, E("BALTOT", 96)))
    check("dashboard total withdrawn = projection cum", close(dash["F9"].value, proj["AB11"].value))
    ref_cum_tax_y8 = sum(t["standalone"] for t in taxes[:8])
    check("dashboard tax-to-date matches reference", close(dash["J9"].value, ref_cum_tax_y8))
    check("dashboard longevity string present",
          isinstance(dash["J15"].value, str) and ("Lasts" in dash["J15"].value or "Depletes" in dash["J15"].value))
    check("scenario table ending balances present",
          all(isinstance(dash.cell(row=22, column=c).value, (int, float)) for c in (3, 4, 5)))

    # --- no error cells anywhere (NA() allowed only on ChartData)
    bad = []
    for ws in wb.worksheets:
        allow_na = ws.title == "ChartData"
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("#"):
                    if allow_na and v == "#N/A":
                        continue
                    bad.append(f"{ws.title}!{cell.coordinate}={v}")
    check("no error cells outside ChartData", not bad, "; ".join(bad[:8]))

    print()
    if fails:
        print(f"{len(fails)} FAILURES:\n  " + "\n  ".join(fails))
        sys.exit(1)
    print("all checks passed ✅")


if __name__ == "__main__":
    main()
