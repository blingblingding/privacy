#!/usr/bin/env python3
"""
Edge-case tests (spec §7 "graceful behaviour on blank/zero entries"):
mutates the Inputs tab, recalculates with LibreOffice, and compares against
the reference model configured identically.
"""

import os, sys, tempfile
import openpyxl
from verify_workbook import recalc, check, close, fails, ROW0
from build_workbook import EL
from reference_model import Inputs, simulate, annual_tax

SRC = os.path.join(os.path.dirname(__file__), "Portfolio_Planner_CA.xlsx")


def mutate(changes):
    wb = openpyxl.load_workbook(SRC)
    ws = wb["Inputs"]
    for addr, v in changes.items():
        ws[addr] = v
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path


def E(wb, colname, month, sheet="Engine (Monthly)"):
    return wb[sheet][f"{EL[colname]}{month + ROW0}"].value


def scenario(title, changes, inp, months_to_check, years_to_check):
    print(f"--- {title}")
    path = mutate(changes)
    wb = recalc(path)
    ref = simulate(inp)
    taxes = annual_tax(ref, inp)
    for t in months_to_check:
        check(f"{title}: balance month {t}", close(E(wb, "BALTOT", t), ref[t - 1]["baltot"]),
              f"xlsx={E(wb, 'BALTOT', t)} ref={ref[t - 1]['baltot']}")
        check(f"{title}: drawn month {t}", close(E(wb, "DRAWTOT", t), ref[t - 1]["drawtot"]),
              f"xlsx={E(wb, 'DRAWTOT', t)} ref={ref[t - 1]['drawtot']}")
    tax = wb["Tax Estimate"]
    for y in years_to_check:
        check(f"{title}: standalone tax year {y}",
              close(tax[f"S{y+3}"].value, taxes[y - 1]["standalone"]),
              f"xlsx={tax[f'S{y+3}'].value} ref={taxes[y - 1]['standalone']}")
    bad = []
    for ws in wb.worksheets:
        if ws.title == "ChartData":
            continue
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    bad.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    check(f"{title}: no error cells", not bad, "; ".join(bad[:6]))
    os.unlink(path)
    return wb


# 1. zero withdrawal — everything must still work, longevity says "Lasts"
wb = scenario(
    "zero withdrawal", {"B28": 0},
    Inputs(wd_amt=0), [12, 240], [1, 20])
check("zero withdrawal: longevity = lasts", "Lasts" in (wb["Dashboard"]["J15"].value or ""))
check("zero withdrawal: withdrawals summary is zero", (wb["Projection"]["AB11"].value or 0) == 0)

# 2. months-unit horizon (18 months), annual withdrawal from month 6, DRIP off,
#    contributions routed 100% to one sub-section, monthly contributions
wb = scenario(
    "18-month horizon / annual wd / DRIP off / routed contribs",
    {"B5": 18, "C5": "Months", "B21": "Monthly", "B23": "", "B25": "Canadian Dividend Stocks",
     "B29": "Annual", "B30": 6, "B42": "No"},
    Inputs(horizon_m=18, contrib_freq="Monthly", contrib_end=0, contrib_dest="Canadian Dividend Stocks",
           wd_freq="Annual", wd_start=6, drip=False),
    [5, 6, 7, 17, 18], [1, 2])
check("annual wd: month 6 and 18 only",
      (E(wb, "DRAWTOT", 6) or 0) > 0 and (E(wb, "DRAWTOT", 7) or 0) < 1e-9 and (E(wb, "DRAWTOT", 18) or 0) > 0)

# 3. bi-weekly withdrawals, funding order reversed (RRSP first), no indexing
wb = scenario(
    "bi-weekly wd / RRSP-first order / no indexing",
    {"B29": "Bi-weekly", "B32": "No", "B33": "RRSP/RRIF", "B34": "TFSA", "B35": "Non-Registered"},
    Inputs(wd_freq="Bi-weekly", wd_indexed=False, order=("RR", "TF", "NR")),
    [61, 120, 240], [6, 20])
check("RRSP-first: RRSP is drawn in month 61",
      (E(wb, "DRAW_RR", 61) or 0) > 0 and (E(wb, "DRAW_NR", 61) or 0) < 1e-9)

print()
if fails:
    print(f"{len(fails)} FAILURES")
    sys.exit(1)
print("all edge-case checks passed ✅")
